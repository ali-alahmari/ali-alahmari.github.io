from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
F='Sakkal Majalla'
d=Document()
cp=d.core_properties; cp.author='علي الأحمري'; cp.last_modified_by='علي الأحمري'; cp.title='دليل رئيس الفريق: نموذج تقويم المعايير للمراجعة'
sec=d.sections[0]; sec.page_width=Cm(21); sec.page_height=Cm(29.7)
for a in ('left_margin','right_margin'): setattr(sec,a,Cm(2.2))
sec.top_margin=Cm(1.8); sec.bottom_margin=Cm(1.8)
def font(r,size,bold=False,color=None):
    r.font.name=F; r.font.size=Pt(size); r.font.bold=bold
    rpr=r._element.get_or_add_rPr(); rf=rpr.find(qn('w:rFonts'))
    if rf is None: rf=OxmlElement('w:rFonts'); rpr.append(rf)
    for k in ('w:ascii','w:hAnsi','w:cs','w:eastAsia'): rf.set(qn(k),F)
    if bold:
        b=OxmlElement('w:bCs'); rpr.append(b)
    cs=OxmlElement('w:szCs'); cs.set(qn('w:val'),str(int(size*2))); rpr.append(cs)
    if color: r.font.color.rgb=RGBColor.from_string(color)
def para(text,size=14,bold=False,color=None,align='r',after=4,before=0,ltr=False):
    p=d.add_paragraph(); pf=p.paragraph_format; pf.space_after=Pt(after); pf.space_before=Pt(before)
    ppr=p._element.get_or_add_pPr(); bd=OxmlElement('w:bidi'); ppr.append(bd)
    p.alignment=WD_ALIGN_PARAGRAPH.RIGHT if align=='r' else WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run(text); font(r,size,bold,color)
    if ltr:
        rpr=r._element.get_or_add_rPr(); rt=OxmlElement('w:rtl'); rt.set(qn('w:val'),'0'); rpr.append(rt)
    return p
G='0D6B63'
para('دليل رئيس الفريق',24,True,G,'c',2)
para('نموذج تقويم المعايير للمراجعة المكتبية ثم الميدانية',16,False,'444444','c',2)
para('إعداد المهندس/ علي الأحمري',14,True,'8a6d1d','c',0)
para('جميع حقوق الملكية الفكرية محفوظة — لا يجوز التعديل أو إعادة النشر دون إذن كتابي',12,False,'555555','c',10)

para('ما الذي يحدث لبياناتك؟',16,True,G,'r',2,6)
para('كل شيء يُحفظ في حسابك أنت في Google: جدول بيانات الفريق ومجلد شواهد المنشآت. لا يمر شيء عبر حساب المُعِدّ ولا يراه. النموذج نفسه ملف جاهز لا يحتاج تعديلًا، ويظهر في أعلى صفحته اسم المُعِدّ وحقوق الملكية، ولا يعمل إن عُدِّل الملف.',14,after=6)

para('أولًا: رابط النموذج',16,True,G,'r',2,6)
para('افتح النموذج من هذا الرابط، وهو نفسه لكل رؤساء الفرق ولا يحتاج استضافة ولا تنزيل:',14,after=2)
para('https://ali-alahmari.github.io/forms/22m-share/',14,True,'0D6B63','r',2,ltr=True)
para('افتحه من Chrome، ويفضَّل أن تحفظه في المفضلة. لا تُرسل هذا الرابط للأعضاء، فلكل عضو رابط شخصي يظهر لك بعد الربط.',14,after=6)

para('ثانيًا: ربط الجدول في حسابك (نحو 5 دقائق)',16,True,G,'r',2,6)
para('1) افتح رابط النموذج، واضغط «مشاركة الفريق». اكتب رمزًا سريًا للفريق إن أردت.',14,after=2)
para('2) افتح script.new بحسابك، واحذف الكود، والصق الكود الظاهر في النافذة (زر «نسخ الكود»)، ثم احفظ.',14,after=2)
para('3) اختر الدالة setup من أعلى المحرر واضغط «تشغيل»، ووافق على الصلاحيات. تنشئ الدالة جدول Google Sheet في حسابك.',14,after=2)
para('4) اضغط «نشر» ثم «عملية نشر جديدة»، النوع «تطبيق ويب»، تنفيذ باسم «أنا»، الوصول «أي شخص». انسخ الرابط المنتهي بـ exec.',14,after=2)
para('5) الصقه في النموذج واضغط «اختبار وربط». تظهر بعدها روابط الأعضاء. أرسل لكل عضو رابطه الشخصي، واحتفظ برابط الرئيس لنفسك.',14,after=6)

para('ثالثًا: بوابة الشواهد الخاصة بك (اختياري، نحو 5 دقائق)',16,True,G,'r',2,6)
para('1) في النموذج افتح تبويب «طلب الشواهد من المنشأة» ثم «إنشاء بوابة الشواهد الخاصة بي». انسخ كود البوابة.',14,after=2)
para('2) افتح script.new جديدًا، والصق الكود، ثم احفظ، ونفّذ الدالة setup ووافق على الصلاحيات.',14,after=2)
para('3) افتح سجل التنفيذ في أسفل المحرر وانسخ «كلمة سر الفريق في البوابة». تنشئ الدالة أيضًا مجلد «بوابة شواهد المنشآت» في Drive الخاص بك.',14,after=2)
para('4) اضغط «نشر» ثم «عملية نشر جديدة»، «تطبيق ويب»، «أنا»، «أي شخص»، وانسخ الرابط المنتهي بـ exec.',14,after=2)
para('5) في النموذج افتح «إعداد الربط مع البوابة»، والصق الرابط وكلمة السر، واضغط «حفظ على هذا الجهاز». بعدها يمكنك نشر الطلب على المنشأة وإرسال رابطها.',14,after=6)

para('عند كل زيارة جديدة',16,True,G,'r',2,6)
para('من رابط الرئيس اضغط «بدء زيارة جديدة» في النموذج، فتُؤرشف الزيارة السابقة ويُنشأ رابط جديد للفريق. وللبوابة زر «زيارة جديدة» خاص بها بمنشآتها وشواهدها وموعدها.',14,after=6)

para('الروابط المستخدمة',16,True,G,'r',2,6)
for t in ['رابط النموذج: ali-alahmari.github.io/forms/22m-share','محرر Apps Script جديد: script.new']:
    para(t,13,after=1,ltr=False)
para('ملاحظة: صفحات المنشأة والفريق في البوابة تُفتح من صفحة جاهزة مستضافة، وتُحفظ بياناتها في Drive الخاص بك وحدك.',13,False,'555555','r',0,6)
d.save('/home/claude/work/site/forms/22m-share/دليل-رئيس-الفريق.docx')
print('ok')
